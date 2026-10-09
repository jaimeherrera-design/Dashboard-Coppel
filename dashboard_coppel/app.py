from pathlib import Path
from html import escape
from io import BytesIO
from base64 import b64encode
from zipfile import BadZipFile
import sqlite3

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from data_cache import load_summary, iter_summary, SUMMARY_KEYS
from report_view import PLOT_CONFIG, render_report
from detail_view import DETAIL_VERSION, numeric_code, phone_prefix, load_detail_masters, classify_detail, combine_groups, render_detail
from outcome_view import render_outcomes
from contact_analysis import render_contact_analysis
from drive_sources import discover_drive_sources, DRIVE_FOLDER_ID
from streamlit.errors import StreamlitSecretNotFoundError


st.set_page_config(
    page_title="Coppel | Contact Center Intelligence",
    page_icon="C",
    layout="wide",
    initial_sidebar_state="expanded",
)

DATA_ROOT = Path(__file__).resolve().parent.parent
CONTACT_MAPPING_FILE = DATA_ROOT / "Mae_contacto.xlsx"
DID_MAPPING_FILE = DATA_ROOT / "Mae_did.xlsx"
LADA_MAPPING_FILE = DATA_ROOT / "Mae_lada.xlsx"
LOGO_FILE = Path(__file__).resolve().parent / "assets" / "coppel-logo.png"
FileSignature = tuple[str, int, int]
CHUNK_SIZE = 250_000
REQUIRED_COLUMNS = [
    "Call end",
    "Talk Time",
    "Wait Time",
    "Wrap up time",
    "Call Type",
    "Campaign Name",
    "Term Reason",
    "Call Outcome name",
]
EXCLUDED_CAMPAIGN_PATTERN = r"test|prueba"
FILTER_VERSION = "campaign-exclusion-v8-persistent-source-summaries"
GROUP_METRICS = ["calls", "contacts", "talk_sum", "wrap_sum", "wait_sum"]
PPT_NAVY = RGBColor(244, 247, 251)
PPT_SURFACE = RGBColor(255, 255, 255)
PPT_LINE = RGBColor(218, 228, 238)
PPT_WHITE = RGBColor(20, 55, 95)
PPT_MUTED = RGBColor(91, 112, 135)
PPT_TEAL = RGBColor(0, 86, 166)
PPT_GOLD = RGBColor(246, 196, 0)


def style_dashboard() -> None:
    st.markdown(
        """
        <style>
        @import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Space+Grotesk:wght@500;600;700&display=swap');
        :root { --ink:#14375f; --muted:#61758e; --teal:#0056a6; --mint:#e8f1fa; --gold:#f6c400; --surface:#ffffff; --surface-2:#f5f8fc; --line:#dbe6f0; }
        html, body, [class*="css"] { font-family:'DM Sans', sans-serif; color:var(--ink); }
        [data-testid="stAppViewContainer"] { background:#f4f7fb; }
        [data-testid="stHeader"] { background:rgba(244,247,251,.94); }
        [data-testid="stSidebar"] { background:#ffffff; border-right:1px solid var(--line); }
        [data-testid="stSidebar"] * { color:#14375f !important; }
        [data-testid="stSidebar"] .stMultiSelect span { color:#14375f !important; }
        h1, h2, h3 { font-family:'Space Grotesk', sans-serif; letter-spacing:0; }
        h1 { font-size:2.2rem; margin-bottom:.15rem; }
        [data-testid="stMetric"] { min-height:116px; box-sizing:border-box; background:var(--surface); border:1px solid var(--line); border-top:4px solid var(--teal); border-radius:12px; padding:17px 20px 16px; box-shadow:0 8px 24px rgba(20,55,95,.07); transition:transform .18s ease, box-shadow .18s ease, border-color .18s ease; overflow:visible; }
        [data-testid="stMetric"]:hover { transform:translateY(-3px); border-color:#9fc1df; box-shadow:0 14px 30px rgba(20,55,95,.14); }
        [data-testid="stMetricLabel"] { color:var(--muted); font-size:.72rem; line-height:1.25; min-height:2.1em; white-space:normal; text-transform:uppercase; letter-spacing:.06em; overflow:visible; }
        [data-testid="stMetricValue"] { color:var(--ink); font-family:'Space Grotesk', sans-serif; font-size:clamp(1.25rem, 2.1vw, 2rem); line-height:1.15; white-space:normal; overflow:visible; }
        [data-testid="stHorizontalBlock"] > [data-testid="column"]:nth-child(2) [data-testid="stMetric"] { border-top-color:var(--gold); }
        [data-testid="stHorizontalBlock"] > [data-testid="column"]:nth-child(4) [data-testid="stMetric"] { border-top-color:var(--gold); }
        .kpi-card { position:relative; min-height:116px; box-sizing:border-box; overflow:hidden; background:#fff; border:1px solid var(--line); border-top:4px solid var(--teal); border-radius:12px; padding:17px 76px 16px 20px; box-shadow:0 8px 24px rgba(20,55,95,.07); transition:transform .18s ease, box-shadow .18s ease, border-color .18s ease; }
        .kpi-card:hover { transform:translateY(-3px); border-color:#9fc1df; box-shadow:0 14px 30px rgba(20,55,95,.14); }
        .kpi-card.gold { border-top-color:var(--gold); }
        .kpi-label { min-height:2.1em; color:var(--muted); font-size:calc(.72rem - 2px); line-height:1.25; text-transform:uppercase; letter-spacing:.06em; }
        .kpi-value { color:var(--ink); font-family:'Space Grotesk', sans-serif; font-size:clamp(1.125rem, 2.1vw - 2px, 1.875rem); font-weight:700; line-height:1.15; white-space:normal; overflow-wrap:anywhere; }
        .kpi-icon { position:absolute; right:18px; bottom:15px; width:34px; height:34px; color:#111; opacity:.9; }
        .kpi-icon svg { display:block; width:100%; height:100%; }
        .eyebrow { color:var(--teal); text-transform:uppercase; letter-spacing:.13em; font-size:.72rem; font-weight:700; }
        .subtitle { color:var(--muted); margin-top:0; }
        .section-title { margin:26px 0 8px; font-family:'Space Grotesk', sans-serif; font-size:1.15rem; font-weight:600; }
        .insight { background:#e9f2fb; border-left:4px solid var(--teal); border-radius:6px; padding:13px 16px; color:#214d78; }
        .insight-grid { display:grid; grid-template-columns:repeat(3, 1fr); gap:14px; margin:10px 0 28px; }
        .insight-card { min-height:190px; background:var(--surface); border:1px solid var(--line); border-top:4px solid #6d7d84; border-radius:10px; padding:18px; box-shadow:0 8px 24px rgba(20,55,95,.07); }
        .insight-card.red { border-top-color:#d9574f; }
        .insight-card.amber { border-top-color:#d99a2b; }
        .insight-card.green { border-top-color:#2d9d78; }
        .insight-card-title { display:flex; align-items:center; gap:8px; color:var(--ink); font-family:'Space Grotesk', sans-serif; font-size:1rem; font-weight:700; margin-bottom:14px; }
        .traffic-light { width:11px; height:11px; border-radius:50%; background:#6d7d84; display:inline-block; }
        .red .traffic-light { background:#d9574f; box-shadow:0 0 0 4px #fbe8e6; }
        .amber .traffic-light { background:#d99a2b; box-shadow:0 0 0 4px #fff3d9; }
        .green .traffic-light { background:#2d9d78; box-shadow:0 0 0 4px #e3f4ed; }
        .insight-card-finding, .insight-card-action { color:#61758e; font-size:.9rem; line-height:1.45; margin-top:10px; }
        .insight-card-finding strong, .insight-card-action strong { color:var(--ink); }
        @media (max-width: 900px) { .insight-grid { grid-template-columns:1fr; } }
        [data-testid="stAppViewContainer"] { background:linear-gradient(180deg,#edf3fb 0,#f5f7fb 460px); }
        [data-testid="stHeader"] { background:rgba(237,243,251,.95); }
        .block-container { padding-top:2rem; max-width:1600px; }
        .executive-banner { position:relative; isolation:isolate; overflow:hidden; display:flex; align-items:center; gap:28px; padding:32px; margin-bottom:24px; border-radius:20px; background:linear-gradient(110deg,#104284 0%,#163f83 55%,#23458f 100%); color:#fff; box-shadow:0 14px 36px rgba(20,55,95,.18); border-bottom:4px solid #f6c400; }
        .executive-banner::after { content:""; position:absolute; z-index:-1; right:0; top:-160px; width:320px; height:420px; border:60px solid rgba(255,255,255,.04); border-radius:50%; pointer-events:none; }
        .banner-logo { width:180px; height:auto; flex-shrink:0; border-radius:10px; }
        .banner-copy { flex:1; min-width:0; border-left:1px solid rgba(255,255,255,.2); padding-left:28px; }
        .banner-eyebrow { font-size:11px; font-weight:700; letter-spacing:.18em; text-transform:uppercase; color:#f6d644; margin-bottom:8px; }
        .executive-banner h1 { font-family:'Space Grotesk',sans-serif; color:#fff; font-size:clamp(1.6rem,2.4vw,2.4rem); line-height:1.15; padding:0; margin:0 0 10px; }
        .banner-subtitle { color:#d6e4fa; font-size:13px; line-height:1.5; }
        .banner-period { flex-shrink:0; padding:15px 20px; background:rgba(255,255,255,.09); border:1px solid rgba(255,255,255,.18); border-radius:12px; }
        .banner-period-label { display:block; font-size:10px; letter-spacing:.12em; text-transform:uppercase; color:#c5d8f4; margin-bottom:6px; }
        .banner-period-value { font-size:13px; font-weight:600; color:#fff; white-space:nowrap; }
        .kpi-card { min-height:120px; border:1px solid #dfe7f2; border-top:4px solid #23458f; border-radius:14px; box-shadow:0 6px 20px rgba(20,55,95,.06); padding:20px 66px 20px 20px; }
        .kpi-card.gold { border-top-color:#f6c400; }
        .kpi-card:hover { transform:translateY(-2px); box-shadow:0 12px 28px rgba(20,55,95,.12); }
        .kpi-label { color:#61758e; text-transform:uppercase; letter-spacing:.06em; font-size:11px; }
        .kpi-value { color:#14375f; font-size:clamp(1.3rem,1.9vw,1.9rem); }
        .kpi-icon { width:38px; height:38px; padding:8px; box-sizing:border-box; border-radius:11px; background:#edf3fc; color:#23458f; opacity:1; }
        .kpi-card.gold .kpi-icon { background:#fff7d7; }
        .st-key-report-content { margin-top:24px; margin-bottom:0; }
        .section-title.dimension-title { margin-top:0; }
        .st-key-detail-content { margin-top:24px; gap:32px; }
        .st-key-outcome-content { margin-top:24px; }
        .st-key-analysis-content { margin-top:24px; gap:24px; }
        .st-key-report-content [class*="st-key-report-"] { margin-bottom:32px; border:0; box-shadow:none; }
        .st-key-report-content [data-testid="stVerticalBlockBorderWrapper"], .st-key-report-content [data-testid="stVerticalBlock"][style*="border:"] { border:0 !important; box-shadow:none; margin-bottom:32px; }
        .report-table-scroll { border-radius:9px; }
        .section-title { color:#23458f; padding-left:12px; border-left:3px solid #f6c400; margin-top:40px; margin-bottom:20px; }
        .report-table-scroll { max-height:330px; overflow:auto; min-height:220px; }
        .report-table { width:100%; border-collapse:collapse; font-family:Arial,sans-serif; font-size:11px; color:#515967; }
        .report-table th, .report-table td { padding:7px 8px; border-bottom:1px solid #f0f1f5; white-space:nowrap; text-align:right; }
        .report-table th:first-child, .report-table td:first-child { text-align:left; }
        .report-table thead th { background:#23458f; color:#fff; position:sticky; top:0; z-index:1; font-weight:600; }
        .report-table tfoot td, .report-table tfoot th { background:#23458f; color:#fff; font-weight:700; position:sticky; bottom:0; }
        .report-panel-title { text-align:center; color:#23458f; font-size:12px; font-weight:700; letter-spacing:.06em; padding:8px 0 16px; }
        .report-heat-scroll { overflow:auto; }
        .report-heat { font-size:10px; }
        .report-heat th, .report-heat td { padding:5px 6px; text-align:center; border:1px solid #fff; }
        .report-heat tbody th { background:#f3f5f8; }
        .report-heat .no-data { background:#f5f6f8; color:#9298a1; }
        @media (max-width: 1100px) { .executive-banner { flex-wrap:wrap; gap:20px; padding:24px; } .banner-period { margin-left:auto; } .banner-logo { width:140px; } }
        @media (max-width: 640px) { .executive-banner { gap:16px; padding:20px; border-radius:14px; } .banner-copy { flex-basis:100%; border-left:0; padding-left:0; } .banner-period { margin-left:0; } .banner-logo { width:120px; } .kpi-card { padding:16px 52px 16px 16px; } .report-table { font-size:10px; } }
        @media (prefers-reduced-motion: reduce) { .kpi-card { transition:none; } .kpi-card:hover { transform:none; } }
        </style>
        """,
        unsafe_allow_html=True,
    )


def render_executive_banner(filters: dict) -> None:
    try:
        logo = b64encode(LOGO_FILE.read_bytes()).decode("ascii")
    except OSError as exc:
        st.error(f"No se pudo cargar el logo del encabezado: {exc}")
        st.stop()
    period = f'{filters["start"].strftime("%d/%m/%Y")} — {filters["end"].strftime("%d/%m/%Y")}'
    st.markdown(
        f'<header class="executive-banner">'
        f'<img class="banner-logo" src="data:image/png;base64,{logo}" alt="Coppel">'
        '<div class="banner-copy"><div class="banner-eyebrow">Contact Center Intelligence</div>'
        '<h1>REPORTE DE GESTION Y TELEFONIA COPPEL</h1>'
        '<div class="banner-subtitle">Visión ejecutiva del rendimiento operativo y la contactabilidad</div></div>'
        '<div class="banner-period"><span class="banner-period-label">Periodo de análisis</span>'
        f'<span class="banner-period-value">{escape(period)}</span></div></header>',
        unsafe_allow_html=True,
    )


def file_signature(file_path: Path) -> FileSignature:
    info = file_path.stat()
    return str(file_path.resolve()), info.st_mtime_ns, info.st_size


def discover_csv_files(root: Path) -> tuple[FileSignature, ...]:
    files = sorted(
        (path for path in root.iterdir() if path.is_file() and path.suffix.casefold() == ".csv"),
        key=lambda path: path.name.casefold(),
    )
    if not files:
        raise ValueError(f"No se encontraron archivos CSV en la raíz: {root}")
    return tuple(file_signature(path) for path in files)


def discover_dashboard_files(root: Path) -> tuple[FileSignature, ...]:
    sources = sorted(
        (path for path in root.iterdir() if path.is_file() and path.suffix.casefold() == ".parquet"),
        key=lambda path: path.name.casefold(),
    )
    if not sources:
        raise ValueError(f"No se encontraron archivos Parquet en la raíz: {root}")
    return tuple(file_signature(path) for path in sources)


def read_file_chunks(path: Path, columns: list[str], *, detail: bool):
    batch_size = min(CHUNK_SIZE, 50_000) if detail else CHUNK_SIZE
    if path.suffix.casefold() == ".parquet":
        with pq.ParquetFile(path) as reader:
            missing = set(columns).difference(reader.schema_arrow.names)
            if missing:
                raise ValueError(f"Faltan columnas requeridas: {', '.join(sorted(missing))}")
            for batch in reader.iter_batches(batch_size=batch_size, columns=columns):
                yield batch.to_pandas()
    elif path.suffix.casefold() == ".csv":
        with pd.read_csv(
            path, sep=";", usecols=columns,
            dtype={"DDI": str, "Phone": str} if detail else None,
            chunksize=batch_size, low_memory=False,
        ) as reader:
            yield from reader
    else:
        raise ValueError(f"Formato de fuente no soportado: {path.suffix}")


def read_source_chunks(files: tuple[FileSignature, ...], status=None, *, detail: bool = False):
    for file_index, (file_path, modified_ns, size) in enumerate(files, start=1):
        path = Path(file_path)
        if file_signature(path) != (file_path, modified_ns, size):
            raise ValueError(f"La fuente cambió durante la carga: {path.name}. Actualiza las fuentes.")
        try:
            columns = REQUIRED_COLUMNS + ["DDI", "Phone"] if detail else REQUIRED_COLUMNS
            for batch_number, chunk in enumerate(read_file_chunks(path, columns, detail=detail), start=1):
                if status is not None and (batch_number == 1 or batch_number % 10 == 0):
                    status.update(
                        label=f"Leyendo {path.name} ({file_index}/{len(files)}): lote {batch_number:,}...",
                        state="running",
                    )
                yield chunk
        except (OSError, ValueError, pa.ArrowException) as exc:
            raise ValueError(f"No se pudo leer la fuente {path.name}: {exc}") from exc
        if file_signature(path) != (file_path, modified_ns, size):
            raise ValueError(f"La fuente cambió durante la carga: {path.name}. Actualiza las fuentes.")


def scan_metadata(files: tuple[FileSignature, ...], status=None) -> dict:
    options = {name: set() for name in ["Call Type", "Campaign Name", "Term Reason"]}
    minimum = None
    maximum = None
    for chunk in read_source_summaries(files, status):
        chunk = chunk.loc[
            ~chunk["Campaign Name"].fillna("").astype(str).str.contains(
                EXCLUDED_CAMPAIGN_PATTERN, case=False, na=False
            )
        ]
        dates = pd.to_datetime(chunk["Call end"], errors="coerce")
        if dates.notna().any():
            minimum = dates.min() if minimum is None else min(minimum, dates.min())
            maximum = dates.max() if maximum is None else max(maximum, dates.max())
        for name in options:
            options[name].update(chunk[name].unique())
    if status is not None:
        status.update(label="Metadatos listos", state="complete")
    return {"options": {key: sorted(value) for key, value in options.items()}, "minimum": minimum, "maximum": maximum}


@st.cache_data(show_spinner="Cargando clasificación de Contacto/No Contacto...")
def load_contact_outcomes(mapping_signature: FileSignature) -> set[str]:
    file_path, _, _ = mapping_signature
    mapping = pd.read_excel(file_path)
    required = ["Call Outcome name", "Contacto / No Contacto"]
    missing = set(required).difference(mapping.columns)
    if missing:
        raise ValueError(f"El maestro de contactos no tiene las columnas requeridas: {', '.join(sorted(missing))}")
    mapping = mapping[required].dropna(how="all").copy()
    mapping.columns = ["call_outcome_name", "contact_flag"]
    mapping["call_outcome_name"] = mapping["call_outcome_name"].fillna("").astype(str).str.strip()
    mapping["contact_flag"] = mapping["contact_flag"].fillna("").astype(str).str.strip().str.casefold()
    mapping["call_outcome_name"] = mapping["call_outcome_name"].str.casefold()
    if mapping.empty or mapping["call_outcome_name"].eq("").any():
        raise ValueError("El maestro de contactos está vacío o contiene resultados de llamada sin nombre.")
    if not mapping["contact_flag"].isin(["contacto", "no contacto"]).all():
        raise ValueError("El maestro debe clasificar cada resultado como Contacto o No Contacto.")
    conflicts = mapping.groupby("call_outcome_name")["contact_flag"].nunique()
    if conflicts.gt(1).any():
        raise ValueError("El maestro contiene clasificaciones contradictorias para un mismo resultado de llamada.")
    if file_signature(Path(file_path)) != mapping_signature:
        raise ValueError("El maestro de contactos cambió durante la carga. Actualiza las fuentes.")
    return set(mapping.loc[mapping["contact_flag"].eq("contacto"), "call_outcome_name"])


@st.cache_data(show_spinner=False)
def dashboard_metadata(files: tuple[FileSignature, ...], version: str) -> dict:
    return scan_metadata(files)


def clean_chunk(chunk: pd.DataFrame) -> pd.DataFrame:
    chunk["Call end"] = pd.to_datetime(chunk["Call end"], errors="coerce")
    for name in ["Talk Time", "Wait Time", "Wrap up time"]:
        chunk[name] = pd.to_numeric(chunk[name], errors="coerce").fillna(0)
    chunk["Call Type"] = chunk["Call Type"].fillna("Sin dato").astype(str).str.strip()
    chunk["Campaign Name"] = chunk["Campaign Name"].fillna("Sin dato").astype(str).str.strip()
    chunk["Term Reason"] = chunk["Term Reason"].fillna("Sin dato").astype(str).str.strip()
    chunk["Call Outcome name"] = chunk["Call Outcome name"].fillna("Sin dato").astype(str).str.strip()
    return chunk


def read_source_summaries(files: tuple[FileSignature, ...], status=None):
    for signature in files:
        summary = load_summary(
            signature,
            lambda: (clean_chunk(raw) for raw in read_source_chunks((signature,), status)),
            status,
        )
        summary["Call end"] = pd.to_datetime(summary["Call end"], errors="coerce")
        yield summary


def read_detail_summaries(files: tuple[FileSignature, ...], status=None):
    def chunks(signature):
        for raw in read_source_chunks((signature,), status, detail=True):
            chunk = clean_chunk(raw)
            chunk["DDI"] = chunk["DDI"].fillna("").map(numeric_code)
            chunk["phone_prefix"] = chunk["Phone"].fillna("").map(phone_prefix)
            yield chunk.drop(columns="Phone")

    for signature in files:
        for summary in iter_summary(
            signature, lambda: chunks(signature), status,
            keys=SUMMARY_KEYS + ["DDI", "phone_prefix"],
            version=DETAIL_VERSION, namespace=DETAIL_VERSION,
        ):
            summary["Call end"] = pd.to_datetime(summary["Call end"], errors="coerce")
            yield summary


@st.cache_data(show_spinner=False, max_entries=2)
def aggregate_detail(files: tuple[FileSignature, ...], contact_signature: FileSignature,
                     did_signature: FileSignature, lada_signature: FileSignature,
                     filters: dict, version: str = DETAIL_VERSION, _status=None) -> dict:
    did, ladas = load_detail_masters(did_signature, lada_signature)
    outcomes = load_contact_outcomes(contact_signature)
    dimensions = {
        "did": ["did"], "did_month": ["did", "month", "provider"],
        "did_campaign": ["did", "Campaign Name", "month"],
        "state": ["state"], "state_month": ["state", "month", "provider"],
    }
    groups = {key: pd.DataFrame(columns=group + GROUP_METRICS) for key, group in dimensions.items()}
    unmapped_did = unmapped_state = 0
    for chunk in read_detail_summaries(files, _status):
        if _status is not None:
            _status.update(label="Aplicando filtros y agrupando DID / Estado...", state="running")
        chunk = chunk.loc[selected_mask(chunk, filters)].copy()
        if chunk.empty:
            continue
        chunk["contacts"] = chunk["calls"] * chunk["Call Outcome name"].str.casefold().isin(outcomes).astype(int)
        chunk = classify_detail(chunk, did, ladas)
        unmapped_did += int(chunk.loc[chunk["did"].eq("Sin cruce"), "calls"].sum())
        unmapped_state += int(chunk.loc[chunk["state"].eq("Sin cruce"), "calls"].sum())
        for key, group in dimensions.items():
            grouped = chunk.groupby(group, as_index=False)[GROUP_METRICS].sum()
            groups[key] = grouped if groups[key].empty else combine_groups([groups[key], grouped], group)
    return {
        **groups,
        "unmapped_did": unmapped_did, "unmapped_state": unmapped_state,
    }


def selected_mask(chunk: pd.DataFrame, filters: dict) -> pd.Series:
    dates = chunk["Call end"]
    campaign_allowed = ~chunk["Campaign Name"].str.contains(
        EXCLUDED_CAMPAIGN_PATTERN, case=False, na=False
    )
    mask = (
        dates.notna()
        & dates.ge(pd.Timestamp(filters["start"]))
        & dates.lt(pd.Timestamp(filters["end"]) + pd.Timedelta(days=1))
        & campaign_allowed
    )
    for field in ["Call Type", "Campaign Name", "Term Reason"]:
        values = filters[field]
        if values:
            mask &= chunk[field].isin(values)
    return mask


@st.cache_data(show_spinner=False)
def aggregate_contact_analysis(files: tuple[FileSignature, ...], mapping_signature: FileSignature, filters: dict) -> pd.DataFrame:
    outcomes = load_contact_outcomes(mapping_signature)
    mapping = pd.read_excel(mapping_signature[0], usecols=["Call Outcome name"])
    known = set(mapping["Call Outcome name"].dropna().astype(str).str.strip().str.casefold())
    if file_signature(Path(mapping_signature[0])) != mapping_signature:
        raise ValueError("El maestro de contactos cambió durante la lectura. Recarga el dashboard.")
    parts = []
    dimensions = ["Campaign Name", "Call Type", "Call Outcome name", "mapped"]
    for chunk in read_source_summaries(files):
        chunk = chunk.loc[selected_mask(chunk, filters)].copy()
        if chunk.empty:
            continue
        normalized = chunk["Call Outcome name"].str.casefold()
        chunk["contacts"] = chunk["calls"] * normalized.isin(outcomes).astype(int)
        chunk["mapped"] = normalized.isin(known)
        parts.append(chunk.groupby(dimensions, as_index=False)[["calls", "contacts"]].sum())
    if not parts:
        return pd.DataFrame(columns=dimensions + ["calls", "contacts"])
    return pd.concat(parts).groupby(dimensions, as_index=False)[["calls", "contacts"]].sum()


@st.cache_data(show_spinner=False)
def aggregate_outcomes(files: tuple[FileSignature, ...], mapping_signature: FileSignature, filters: dict) -> pd.DataFrame:
    parts = []
    outcomes = load_contact_outcomes(mapping_signature)
    dimensions = ["Call Outcome name", "month", "date"]
    for chunk in read_source_summaries(files):
        chunk = chunk.loc[selected_mask(chunk, filters)].copy()
        if chunk.empty:
            continue
        chunk["month"] = chunk["Call end"].dt.to_period("M").astype(str)
        chunk["date"] = chunk["Call end"].dt.strftime("%Y-%m-%d")
        chunk["contacts"] = chunk["calls"] * chunk["Call Outcome name"].str.casefold().isin(outcomes).astype(int)
        parts.append(chunk.groupby(dimensions, as_index=False)[GROUP_METRICS].sum())
    return combine_groups(parts, dimensions)


@st.cache_data(show_spinner=False)
def contact_hourly_days(files: tuple[FileSignature, ...], mapping_signature: FileSignature, filters: dict) -> pd.DataFrame:
    parts = []
    outcomes = load_contact_outcomes(mapping_signature)
    for chunk in read_source_summaries(files):
        chunk = chunk.loc[selected_mask(chunk, filters)].copy()
        if chunk.empty:
            continue
        chunk["day_of_month"] = chunk["Call end"].dt.day
        chunk["hour"] = chunk["Call end"].dt.hour
        chunk["contacts"] = chunk["calls"] * chunk["Call Outcome name"].str.casefold().isin(outcomes).astype(int)
        parts.append(chunk.groupby(["day_of_month", "hour"], as_index=False)[["calls", "contacts"]].sum())
    if not parts:
        return pd.DataFrame(columns=["day_of_month", "hour", "calls", "contacts"])
    return pd.concat(parts).groupby(["day_of_month", "hour"], as_index=False)[["calls", "contacts"]].sum()


@st.cache_data(show_spinner="Procesando registros seleccionados...")
def aggregate(files: tuple[FileSignature, ...], mapping_signature: FileSignature, filter_version: str, filters: dict) -> dict:
    total = contacts = talk_sum = wrap_sum = wait_sum = 0.0
    by_month, by_day, by_day_of_month, by_weekday, by_hour = [], [], [], [], []
    by_call_type, by_campaign, by_term = [], [], []
    heat = []
    contact_outcomes = load_contact_outcomes(mapping_signature)
    for chunk in read_source_summaries(files):
        chunk = chunk.loc[selected_mask(chunk, filters)].copy()
        if chunk.empty:
            continue
        chunk["month"] = chunk["Call end"].dt.to_period("M").astype(str)
        chunk["day"] = chunk["Call end"].dt.date.astype(str)
        chunk["day_of_month"] = chunk["Call end"].dt.day
        chunk["hour"] = chunk["Call end"].dt.hour
        chunk["weekday"] = chunk["Call end"].dt.day_name().map({"Monday":"Lun", "Tuesday":"Mar", "Wednesday":"Mie", "Thursday":"Jue", "Friday":"Vie", "Saturday":"Sab", "Sunday":"Dom"})
        chunk["contact"] = chunk["calls"] * chunk["Call Outcome name"].str.casefold().isin(contact_outcomes).astype(int)
        total += chunk["calls"].sum()
        contacts += chunk["contact"].sum()
        talk_sum += chunk["talk_sum"].sum()
        wrap_sum += chunk["wrap_sum"].sum()
        wait_sum += chunk["wait_sum"].sum()
        for target, keys in [(by_month, ["month"]), (by_day, ["day"]), (by_day_of_month, ["day_of_month"]), (by_weekday, ["weekday"]), (by_hour, ["hour"]), (by_call_type, ["Call Type"]), (by_campaign, ["Campaign Name"]), (by_term, ["Term Reason"]), (heat, ["weekday", "hour"])]:
            grouped = chunk.groupby(keys, dropna=False).agg(
                calls=("calls", "sum"),
                contacts=("contact", "sum"),
                talk_sum=("talk_sum", "sum"),
                wrap_sum=("wrap_sum", "sum"),
                wait_sum=("wait_sum", "sum"),
            ).reset_index()
            target.append(grouped)
    def combine(parts: list) -> pd.DataFrame:
        if not parts:
            return pd.DataFrame()
        keys = [column for column in parts[0].columns if column not in GROUP_METRICS]
        return pd.concat(parts).groupby(keys, as_index=False)[GROUP_METRICS].sum()
    return {
        "total": int(total), "contacts": int(contacts), "talk_sum": talk_sum, "wrap_sum": wrap_sum, "wait_sum": wait_sum,
        "month": combine(by_month), "day": combine(by_day), "day_of_month": combine(by_day_of_month), "weekday": combine(by_weekday), "hour": combine(by_hour), "call_type": combine(by_call_type),
        "campaign": combine(by_campaign), "term": combine(by_term), "heat": combine(heat),
    }


def fmt_number(value: float) -> str:
    return f"{value:,.0f}".replace(",", ".")


def render_kpi_card(label: str, value: str, icon: str, gold: bool = False) -> None:
    icons = {
        "calls": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M4 13v-2a8 8 0 0 1 16 0v2"/><rect x="2" y="11" width="5" height="8" rx="2"/><rect x="17" y="11" width="5" height="8" rx="2"/><path d="M20 19v1a2 2 0 0 1-2 2h-4"/></svg>',
        "contacts": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><circle cx="8" cy="7" r="4"/><path d="M2 21v-2a6 6 0 0 1 10-4.5"/><circle cx="17" cy="17" r="5"/><path d="m14.5 17 1.5 1.5 3-3"/></svg>',
        "rate": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><circle cx="12" cy="12" r="10"/><path d="m8 16 8-8"/><circle cx="8" cy="8" r="1.5"/><circle cx="16" cy="16" r="1.5"/></svg>',
        "idle": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><circle cx="12" cy="12" r="10"/><path d="M9 8v8M15 8v8"/></svg>',
        "tmo": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><circle cx="12" cy="14" r="8"/><path d="M9 2h6M12 2v4m6 2 2-2M12 10v4l3 2"/></svg>',
    }
    card_class = "kpi-card gold" if gold else "kpi-card"
    st.markdown(
        f'<div class="{card_class}"><div class="kpi-label">{escape(label)}</div><div class="kpi-value">{escape(value)}</div><div class="kpi-icon">{icons[icon]}</div></div>',
        unsafe_allow_html=True,
    )


def insight_dimension(result: dict, total: int, key: str, label: str) -> pd.DataFrame:
    table = result[key].copy()
    if table.empty:
        return table
    dimension = table.columns[0]
    table["tmo"] = (table["talk_sum"] + table["wrap_sum"]) / table["calls"].replace(0, 1)
    table["contact_rate"] = table["contacts"] / table["calls"].replace(0, 1)
    table["dimension_field"] = dimension
    table["dimension_label"] = label
    table["share"] = table["calls"] / max(total, 1)
    return table


def render_insight_card(status: str, title: str, finding: str, action: str) -> None:
    st.markdown(
        f'''<div class="insight-card {status}">
            <div class="insight-card-title"><span class="traffic-light"></span>{escape(title)}</div>
            <div class="insight-card-finding"><strong>Hallazgo:</strong> {escape(finding)}</div>
            <div class="insight-card-action"><strong>Qué hacer:</strong> {escape(action)}</div>
        </div>''',
        unsafe_allow_html=True,
    )


def style_chart(fig, height: int) -> None:
    fig.update_layout(
        template="plotly_white",
        paper_bgcolor="#ffffff",
        plot_bgcolor="#ffffff",
        font=dict(family="DM Sans, sans-serif", color="#14375f"),
        title_font=dict(family="Space Grotesk, sans-serif", color="#14375f"),
        coloraxis_colorbar=dict(tickfont=dict(color="#61758e")),
        height=height,
    )


def chart_png(fig, width=1400, height=700) -> BytesIO | None:
    try:
        return BytesIO(fig.to_image(format="png", width=width, height=height, scale=1))
    except Exception:
        return None


def transparent_icon(kind: str) -> BytesIO:
    image = Image.new("RGBA", (240, 240), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    color = (0, 86, 166, 255)
    gold = (246, 196, 0, 255)
    if kind == "pulse":
        draw.line([(25, 130), (70, 130), (95, 65), (135, 175), (165, 105), (215, 105)], fill=color, width=12, joint="curve")
        draw.ellipse((82, 52, 108, 78), fill=gold)
    else:
        draw.rounded_rectangle((45, 55, 195, 190), radius=16, outline=color, width=10)
        draw.line([(70, 155), (100, 120), (130, 140), (170, 88)], fill=gold, width=10)
        draw.ellipse((160, 78, 180, 98), fill=gold)
    output = BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    return output


def ppt_text(slide, text, left, top, width, height, size=18, color=PPT_WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(.04)
    frame.margin_right = Inches(.04)
    frame.margin_top = Inches(.03)
    frame.margin_bottom = Inches(.03)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def ppt_background(slide, title, section=None):
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = PPT_NAVY
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(.12)).fill.solid()
    top_line = slide.shapes[-1]
    top_line.fill.fore_color.rgb = PPT_TEAL
    top_line.line.fill.background()
    ppt_text(slide, "COPPEL  /  CONTACT CENTER INTELLIGENCE", .55, .28, 6.5, .25, 9, PPT_TEAL, True)
    ppt_text(slide, title, .55, .62, 10.5, .48, 24, PPT_WHITE, True)
    if section:
        ppt_text(slide, section, 10.0, .68, 2.75, .28, 10, PPT_MUTED, False, PP_ALIGN.RIGHT)


def ppt_card(slide, left, top, width, height, title, value, accent=PPT_TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PPT_SURFACE
    shape.line.color.rgb = PPT_LINE
    shape.line.width = Pt(1)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(.07), Inches(height)).fill.solid()
    accent_shape = slide.shapes[-1]
    accent_shape.fill.fore_color.rgb = accent
    accent_shape.line.fill.background()
    ppt_text(slide, title.upper(), left + .22, top + .18, width - .4, .25, 9, PPT_MUTED, True)
    ppt_text(slide, value, left + .22, top + .55, width - .4, .5, 25, PPT_WHITE, True)


def insight_records(result, total, tmo, contact_rate):
    tables = [
        insight_dimension(result, total, "campaign", "campaña"),
        insight_dimension(result, total, "day_of_month", "día del mes"),
        insight_dimension(result, total, "weekday", "día de semana"),
        insight_dimension(result, total, "hour", "hora"),
    ]
    tables = [table for table in tables if not table.empty]
    all_dimensions = pd.concat(tables, ignore_index=True) if tables else pd.DataFrame()
    heat = result["heat"].copy() if not result["heat"].empty else pd.DataFrame()
    minimum_share = .01
    eligible = all_dimensions.loc[all_dimensions["share"] >= minimum_share] if not all_dimensions.empty else pd.DataFrame()
    high_tmo = eligible.sort_values("tmo", ascending=False) if not eligible.empty else pd.DataFrame()
    high_idle = eligible.assign(idle=lambda frame: frame["wait_sum"] / frame["calls"].replace(0, 1)).sort_values("idle", ascending=False) if not eligible.empty else pd.DataFrame()
    low_contact = eligible.sort_values("contact_rate") if not eligible.empty else pd.DataFrame()
    strongest_hour = result["hour"].sort_values("contacts", ascending=False).head(1) if not result["hour"].empty else pd.DataFrame()
    weekday_insights = insight_dimension(result, total, "weekday", "día de semana")
    best_weekday = weekday_insights.sort_values("contact_rate", ascending=False).head(1) if not weekday_insights.empty else pd.DataFrame()
    best_campaign = insight_dimension(result, total, "campaign", "campaña")
    best_campaign = best_campaign.loc[best_campaign["share"] >= minimum_share].sort_values("contact_rate", ascending=False).head(1) if not best_campaign.empty else pd.DataFrame()
    weak_slot = heat.loc[heat["calls"] >= heat["calls"].sum() * minimum_share].sort_values("contacts").head(1) if not heat.empty else pd.DataFrame()

    cards = []

    if not low_contact.empty and contact_rate:
        row = low_contact.iloc[0]
        value = str(row[row["dimension_field"]])
        cards.append(("red", "Campaña con brecha de contacto", f"{value} registra {row['contact_rate']:.1%} de contacto en {int(row['calls']):,} llamadas, frente al {contact_rate:.1%} general.", f"Priorizar esta campaña en la revisión operativa: moverla a las horas y días con mejor contacto y comparar su script contra la campaña líder antes de aumentar volumen."))
    else:
        cards.append(("red", "Segmento con menor contactabilidad", "No hay un grupo claramente débil por tasa de contacto.", "Mantener un monitoreo semanal por campaña para detectar pérdida de contacto antes de que impacte el objetivo."))

    if not high_tmo.empty and tmo:
        row = high_tmo.iloc[0]
        value = str(row[row["dimension_field"]])
        gap = row["tmo"] / tmo - 1 if tmo else 0
        cards.append(("red", "TMO que limita capacidad", f"{value} tiene TMO de {row['tmo']:.1f} s, {gap:.0%} por encima del promedio, con {int(row['calls']):,} llamadas.", f"Auditar las llamadas de {value}, separar conversación efectiva de cierre administrativo y fijar una meta de reducción de TMO para liberar capacidad de marcación."))
    else:
        cards.append(("red", "Fricción operativa del grupo", "No aparece un segmento con TMO significativamente más alto que el promedio.", "Mantener control de TMO por campaña y motivo de cierre para no perder productividad por desgaste operativo."))

    if not weak_slot.empty:
        row = weak_slot.iloc[0]
        idle_slot = row["wait_sum"] / row["calls"] if row["calls"] else 0
        cards.append(("amber", "Bloque horario a corregir", f"{row['weekday']} a las {int(row['hour']):02d}:00 concentra solo {int(row['contacts']):,} contactos y un Idle de {idle_slot:.1f} s.", f"Reducir la dotación en ese bloque, mover llamadas a la franja de mayor contacto y revisar la disponibilidad de agentes para evitar espera improductiva."))
    else:
        cards.append(("amber", "Horario de peor rendimiento", "No hay un bloque claramente incómodo en la curva horaria.", "Mantener el análisis por hora para reaccionar antes de que falle la ventana principal de contacto."))

    if not best_weekday.empty:
        best_day_value = str(best_weekday.iloc[0]["weekday"])
        cards.append(("amber", "Día de semana para concentrar volumen", f"{best_day_value} alcanza {best_weekday.iloc[0]['contact_rate']:.1%} de contacto sobre {int(best_weekday.iloc[0]['calls']):,} llamadas.", f"Trasladar llamadas de recuperación a {best_day_value}, manteniendo el control de TMO e Idle para que el aumento de volumen no deteriore la capacidad."))
    else:
        cards.append(("amber", "Mecanismo para aumentar contacto", "No hay un día claramente dominante en el rango analizado.", "Usar la mejor ventana de contacto como referencia y ajustar carga por día con base en la evolución semanal."))

    if not strongest_hour.empty:
        strongest_hour_value = int(strongest_hour.iloc[0]["hour"])
        cards.append(("green", "Ventana más efectiva para presionar volumen", f"La hora {strongest_hour_value:02d}:00 es la mejor ventana de contacto del periodo.", f"Aumentar el volumen de prospección y recuperación en las {strongest_hour_value:02d}:00 h y priorizar a los agentes con mejor tasa de cierre en esa franja."))
    else:
        cards.append(("green", "Ventana más efectiva para presionar volumen", "No aparece una ventana claramente superior por contacto.", "Mantener la mejor hora como referencia y reforzar con análisis semanal para sostener la productividad."))

    if not best_campaign.empty:
        campaign_name = str(best_campaign.iloc[0]["Campaign Name"])
        cards.append(("green", "Campaña benchmark", f"{campaign_name} lidera con {best_campaign.iloc[0]['contact_rate']:.1%} de contacto y un TMO de {best_campaign.iloc[0]['tmo']:.1f} s.", f"Replicar su secuencia, argumento y distribución horaria en campañas con brecha, conservando el mismo estándar de TMO."))
    else:
        cards.append(("green", "Patrón a replicar", "No hay una campaña claramente superior en el periodo actual.", "Tomar la campaña con mejor tasa de contacto como benchmark para estandarizar el llamado y el cierre."))

    if not high_idle.empty:
        row = high_idle.iloc[0]
        value = str(row[row["dimension_field"]])
        idle = row["idle"]
        cards.append(("amber", "Idle que debe reducirse", f"{value} presenta el Idle más alto entre las dimensiones relevantes: {idle:.1f} s por llamada.", f"Revisar pausas, disponibilidad y distribución de agentes en {value}; cada segundo recuperado debe convertirse en llamadas adicionales dentro de las mejores ventanas."))

    if not result["day_of_month"].empty:
        day_month = insight_dimension(result, total, "day_of_month", "día del mes")
        day_month = day_month.loc[day_month["share"] >= minimum_share].sort_values("contact_rate", ascending=False).head(1)
        if not day_month.empty:
            row = day_month.iloc[0]
            cards.append(("green", "Día del mes con mejor conversión", f"El día {int(row['day_of_month'])} alcanza {row['contact_rate']:.1%} de contacto, con TMO de {row['tmo']:.1f} s e Idle de {row['wait_sum'] / row['calls']:.1f} s.", f"Reservar capacidad para campañas de mayor valor alrededor del día {int(row['day_of_month'])} y replicar la combinación de horario y dotación que produce este resultado."))

    cards = cards[:8]
    while len(cards) < 8:
        cards.append(("green", "Sostenimiento operativo", "La operación está estable y no requiere acción urgente.", "Mantener seguimiento de campaña, horario y motivo de cierre para proteger la tasa de contacto en la próxima semana."))
    return sorted(cards, key=lambda card: {"red": 0, "amber": 1, "green": 2}[card[0]])


def build_pptx(result, filters, total, contact_rate, idle, tmo) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    period = f"{filters['start'].strftime('%d %b %Y')} - {filters['end'].strftime('%d %b %Y')}"

    slide = prs.slides.add_slide(blank)
    ppt_background(slide, "Contact Center Pulse", "Reporte ejecutivo")
    ppt_text(slide, "Lectura gerencial de llamadas, contacto y eficiencia operativa", .58, 1.55, 7.4, .55, 22, PPT_WHITE, True)
    ppt_text(slide, period, .6, 2.28, 4.5, .3, 13, PPT_MUTED)
    slide.shapes.add_picture(transparent_icon("pulse"), Inches(9.5), Inches(1.45), width=Inches(2.1), height=Inches(2.1))
    ppt_text(slide, "Base analizada sin campañas de prueba o test", .6, 5.95, 6.4, .3, 12, PPT_TEAL, True)

    slide = prs.slides.add_slide(blank)
    ppt_background(slide, "Resumen ejecutivo", period)
    for left, title, value, accent in [(.6, "Total llamadas", fmt_number(total), PPT_TEAL), (3.75, "Contactos", f"{fmt_number(result['contacts'])}  |  {contact_rate:.1%}", PPT_GOLD), (6.9, "Idle time promedio", f"{idle:.1f} s", PPT_TEAL), (10.05, "TMO", f"{tmo:.1f} s", PPT_GOLD)]:
        ppt_card(slide, left, 1.55, 2.7, 1.45, title, value, accent)
    ppt_text(slide, "Lectura rápida", .65, 3.55, 2, .3, 14, PPT_TEAL, True)
    best_day = result["day"].loc[result["day"]["contacts"].idxmax(), "day"] if not result["day"].empty else "sin datos"
    best_hour = int(result["hour"].loc[result["hour"]["contacts"].idxmax(), "hour"]) if not result["hour"].empty else 0
    ppt_text(slide, f"El mejor día de contacto fue {best_day} y la mejor hora fue {best_hour:02d}:00. La contactabilidad general es de {contact_rate:.1%}.", .65, 4.05, 11.7, .75, 20, PPT_WHITE, False)

    monthly = result["month"].sort_values("month").copy()
    monthly["contact rate"] = monthly["contacts"] / monthly["calls"].replace(0, 1)
    monthly_fig = px.bar(monthly, x="month", y="calls", color="contact rate", color_continuous_scale=["#dcecf8", "#0056a6"], labels={"month": "Mes", "calls": "Llamadas", "contact rate": "Contacto"}, title="Volumen mensual y contacto")
    style_chart(monthly_fig, 650)
    hourly = result["hour"].sort_values("hour")
    hourly_fig = px.line(hourly, x="hour", y="contacts", markers=True, labels={"hour": "Hora", "contacts": "Contactos"}, title="Contactos por hora")
    hourly_fig.update_traces(line_color="#f2bd55", marker_color="#f2bd55")
    style_chart(hourly_fig, 650)
    slide = prs.slides.add_slide(blank)
    ppt_background(slide, "Ritmo de contacto", period)
    for left, figure in [(.6, monthly_fig), (6.9, hourly_fig)]:
        image = chart_png(figure)
        if image:
            slide.shapes.add_picture(image, Inches(left), Inches(1.35), width=Inches(5.85), height=Inches(4.45))

    heat = result["heat"].copy()
    order = ["Lun", "Mar", "Mie", "Jue", "Vie", "Sab", "Dom"]
    heat["weekday"] = pd.Categorical(heat["weekday"], categories=order, ordered=True)
    heat = heat.pivot(index="weekday", columns="hour", values="contacts").reindex(order).fillna(0)
    heat_fig = px.imshow(heat, color_continuous_scale=["#edf4fb", "#4c91ca", "#14375f"], labels={"x": "Hora", "y": "Día", "color": "Contactos"}, aspect="auto", text_auto=False, title="Mapa de calor de contacto")
    style_chart(heat_fig, 650)
    top_campaigns = result["campaign"].sort_values("contacts", ascending=False).head(8).sort_values("contacts")
    campaign_fig = px.bar(top_campaigns, x="contacts", y="Campaign Name", orientation="h", labels={"contacts": "Contactos", "Campaign Name": "Campaña"}, title="Campañas con más contactos")
    campaign_fig.update_traces(marker_color="#0056a6")
    style_chart(campaign_fig, 650)
    slide = prs.slides.add_slide(blank)
    ppt_background(slide, "Dónde concentrar la gestión", period)
    for left, figure in [(.6, heat_fig), (6.9, campaign_fig)]:
        image = chart_png(figure)
        if image:
            slide.shapes.add_picture(image, Inches(left), Inches(1.35), width=Inches(5.85), height=Inches(4.45))

    slide = prs.slides.add_slide(blank)
    ppt_background(slide, "Insights y próximos pasos", "Cierre ejecutivo")
    ppt_text(slide, "Prioridades para la siguiente revisión", .62, 1.18, 7.5, .35, 16, PPT_TEAL, True)
    cards = insight_records(result, total, tmo, contact_rate)
    for index, (status, title, finding, action) in enumerate(cards):
        left = .62 + index * 4.22
        accent = {"red": RGBColor(217, 87, 79), "amber": PPT_GOLD, "green": RGBColor(45, 157, 120)}[status]
        ppt_card(slide, left, 1.85, 3.82, 3.8, title, status.upper(), accent)
        ppt_text(slide, "Hallazgo", left + .25, 2.75, 3.2, .25, 10, accent, True)
        ppt_text(slide, finding, left + .25, 3.05, 3.25, .85, 14, PPT_WHITE)
        ppt_text(slide, "Qué hacer", left + .25, 4.15, 3.2, .25, 10, accent, True)
        ppt_text(slide, action, left + .25, 4.45, 3.25, .9, 13, PPT_MUTED)

    output = BytesIO()
    prs.save(output)
    return output.getvalue()


def main() -> None:
    style_dashboard()
    st.sidebar.markdown("## COPPEL\n### Contact Center Intelligence")
    st.sidebar.caption("Modelo de análisis operativo · 2026")
    try:
        try:
            account_info = dict(st.secrets["gcp_service_account"]) if "gcp_service_account" in st.secrets else None
        except StreamlitSecretNotFoundError:
            account_info = None
        with st.status("Consultando Parquet y maestros en Google Drive...", expanded=False) as drive_status:
            sources = discover_drive_sources(DATA_ROOT, drive_status, service_account_info=account_info)
            files = sources.parquet
            drive_status.update(label=f"Drive sincronizado: {len(files)} Parquet y 3 maestros", state="complete")
        mapping_signature = sources.masters["Mae_contacto.xlsx"]
        load_contact_outcomes(mapping_signature)
        metadata_signature = (FILTER_VERSION, files)
        if st.session_state.get("metadata_signature") != metadata_signature:
            with st.spinner("Preparando la base de datos..."):
                metadata = dashboard_metadata(files, FILTER_VERSION)
            st.session_state["source_metadata"] = metadata
            st.session_state["metadata_signature"] = metadata_signature
        else:
            metadata = st.session_state["source_metadata"]
    except (OSError, ValueError, ImportError, BadZipFile, sqlite3.Error, pd.errors.DatabaseError) as exc:
        st.error(f"No se pudieron cargar las fuentes: {exc}")
        st.stop()
    if metadata["minimum"] is None or metadata["maximum"] is None:
        st.warning("Los Parquet no contienen llamadas con fechas válidas fuera de campañas test/prueba.")
        st.stop()
    date_min = metadata["minimum"].date()
    date_max = metadata["maximum"].date()
    date_range = st.sidebar.date_input("Periodo", value=(date_min, date_max), min_value=date_min, max_value=date_max)
    if len(date_range) != 2:
        st.warning("Selecciona una fecha inicial y una fecha final.")
        st.stop()
    filters = {"start": date_range[0], "end": date_range[1]}
    for field, label in [("Call Type", "Call type"), ("Campaign Name", "Campaign name"), ("Term Reason", "Term reason")]:
        filters[field] = st.sidebar.multiselect(label, metadata["options"][field], default=[])
    st.sidebar.divider()
    st.sidebar.caption(
        "Se suman todos los Parquet de la carpeta privada de Google Drive al abrir o recargar. Los archivos de la raíz y las subcarpetas no se incluyen."
    )
    st.sidebar.markdown(f"[Carpeta de fuentes en Drive](https://drive.google.com/drive/folders/{DRIVE_FOLDER_ID})")
    with st.sidebar.expander("Fuentes incluidas"):
        for file_path, _, _ in files:
            st.caption(Path(file_path).name)
        for name in sources.masters:
            st.caption(f"Maestro en Drive: {name}")

    try:
        result = aggregate(files, mapping_signature, FILTER_VERSION, filters)
    except (OSError, ValueError, ImportError, BadZipFile, sqlite3.Error, pd.errors.DatabaseError) as exc:
        st.error(f"No se pudieron procesar las fuentes: {exc}")
        st.stop()
    total = result["total"]
    contact_rate = result["contacts"] / total if total else 0
    idle = result["wait_sum"] / total if total else 0
    tmo = (result["talk_sum"] + result["wrap_sum"]) / total if total else 0
    st.sidebar.divider()
    st.sidebar.markdown("### Exportar reporte")
    if st.sidebar.button("Generar PPT ejecutiva", use_container_width=True):
        with st.spinner("Construyendo presentación..."):
            ppt_data = build_pptx(result, filters, total, contact_rate, idle, tmo)
        st.sidebar.download_button(
            "Descargar PPTX",
            data=ppt_data,
            file_name=f"Coppel_Reporte_Ejecutivo_{filters['end'].strftime('%Y%m%d')}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            use_container_width=True,
        )
    render_executive_banner(filters)
    kpis = st.columns(5)
    with kpis[0]:
        render_kpi_card("Total llamadas", fmt_number(total), "calls")
    with kpis[1]:
        render_kpi_card("Contactos", fmt_number(result["contacts"]), "contacts", True)
    with kpis[2]:
        render_kpi_card("% de contactos", f"{contact_rate:.2%}", "rate")
    with kpis[3]:
        render_kpi_card("Idle time promedio", f"{idle:.1f} s", "idle")
    with kpis[4]:
        render_kpi_card("TMO", f"{tmo:.1f} s", "tmo", True)

    principal, detail, outcomes, analysis = st.tabs(
        ["DASHBOARD PRINCIPAL", "DETALLE DID / ESTADO", "RESULTADO LLAMADAS", "ANÁLISIS CONTACTABILIDAD"], key="dashboard-pages", on_change="rerun",
    )
    if principal.open:
        with principal:
            render_principal(result, files, mapping_signature, filters, total, contact_rate, tmo)
    if detail.open:
        with detail:
            with st.container(border=False, key="detail-content"):
                detail_status = st.status("Preparando detalle DID / Estado...", expanded=False)
                try:
                    did_signature = sources.masters["Mae_did.xlsx"]
                    lada_signature = sources.masters["Mae_lada.xlsx"]
                    detail_result = aggregate_detail(
                        files, mapping_signature, did_signature, lada_signature, filters,
                        _status=detail_status,
                    )
                    render_detail(detail_result, result)
                    detail_status.update(label="Detalle DID / Estado listo", state="complete")
                except (OSError, ValueError, ImportError, BadZipFile, sqlite3.Error, pd.errors.DatabaseError) as exc:
                    detail_status.update(label="No se pudo preparar el detalle DID / Estado", state="error")
                    st.error(f"No se pudo cargar el detalle DID / Estado: {exc}")
    if outcomes.open:
        with outcomes:
            with st.container(border=False, key="outcome-content"):
                try:
                    with st.spinner("Preparando resultados de llamadas..."):
                        outcome_result = aggregate_outcomes(files, mapping_signature, filters)
                    render_outcomes(outcome_result)
                except (OSError, ValueError, ImportError, BadZipFile, sqlite3.Error, pd.errors.DatabaseError) as exc:
                    st.error(f"No se pudieron cargar los resultados de llamadas: {exc}")
    if analysis.open:
        with analysis:
            with st.container(border=False, key="analysis-content"):
                try:
                    with st.spinner("Analizando volumen, clasificación y resultados de llamada..."):
                        analysis_result = aggregate_contact_analysis(files, mapping_signature, filters)
                    render_contact_analysis(analysis_result, result)
                except (OSError, ValueError, ImportError, BadZipFile, sqlite3.Error, pd.errors.DatabaseError) as exc:
                    st.error(f"No se pudo cargar el análisis de contactabilidad: {exc}")


def render_principal(result: dict, files: tuple[FileSignature, ...], mapping_signature: FileSignature,
                     filters: dict, total: float, contact_rate: float, tmo: float) -> None:
    try:
        hourly_days = contact_hourly_days(files, mapping_signature, filters)
    except (OSError, ValueError, ImportError, BadZipFile, sqlite3.Error, pd.errors.DatabaseError) as exc:
        st.error(f"No se pudo cargar el detalle visual por hora: {exc}")
        st.stop()
    with st.container(border=False, key="report-content"):
        render_report(result, hourly_days)

    st.markdown('<div class="section-title dimension-title">Lectura por dimensión</div>', unsafe_allow_html=True)
    campaign_scatter = result["campaign"].copy()
    if not campaign_scatter.empty:
        campaign_scatter["contact_rate"] = campaign_scatter["contacts"] / campaign_scatter["calls"].replace(0, 1)
        campaign_scatter["tmo"] = (campaign_scatter["talk_sum"] + campaign_scatter["wrap_sum"]) / campaign_scatter["calls"].replace(0, 1)
        campaign_scatter["idle"] = campaign_scatter["wait_sum"] / campaign_scatter["calls"].replace(0, 1)
        campaign_scatter["campaign"] = campaign_scatter["Campaign Name"].astype(str)
        campaign_scatter = campaign_scatter.loc[campaign_scatter["calls"] >= campaign_scatter["calls"].sum() * .01].copy()
        if not campaign_scatter.empty:
            contact_q25, contact_q75 = campaign_scatter["contact_rate"].quantile([.25, .75])
            tmo_q25, tmo_q75 = campaign_scatter["tmo"].quantile([.25, .75])

            def classify_campaign(row):
                high_contact = row["contact_rate"] >= contact_q75
                low_contact = row["contact_rate"] <= contact_q25
                low_tmo = row["tmo"] <= tmo_q25
                high_tmo = row["tmo"] >= tmo_q75
                if high_contact and low_tmo:
                    return "Benchmark: alto contacto / bajo TMO"
                if high_contact and high_tmo:
                    return "Potencial: alto contacto / alto TMO"
                if low_contact and high_tmo:
                    return "Crítica: bajo contacto / alto TMO"
                if low_contact and low_tmo:
                    return "Revisar: bajo contacto / bajo TMO"
                return "Intermedia"

            campaign_scatter["segment"] = campaign_scatter.apply(classify_campaign, axis=1)
            segment_colors = {
                "Benchmark: alto contacto / bajo TMO": "#2d9d78",
                "Potencial: alto contacto / alto TMO": "#f6c400",
                "Crítica: bajo contacto / alto TMO": "#d9574f",
                "Revisar: bajo contacto / bajo TMO": "#d99a2b",
                "Intermedia": "#6d7d84",
            }
            fig = px.scatter(
                campaign_scatter,
                x="tmo",
                y="contact_rate",
                size="calls",
                color="segment",
                hover_name="campaign",
                hover_data={"segment":False, "tmo":":.1f", "contact_rate":":.1%", "calls":":,.0f", "idle":":.1f"},
                color_discrete_map=segment_colors,
                labels={"tmo":"TMO (segundos)", "contact_rate":"Contactabilidad", "calls":"Llamadas", "idle":"Idle (segundos)", "segment":"Clasificación"},
                title="TMO y contactabilidad por campaña: cuartiles y zonas de acción",
            )
            fig.update_traces(marker=dict(line=dict(width=1, color="#ffffff"), opacity=.85))
            fig.add_vline(x=tmo, line_dash="dash", line_color="#0056a6", annotation_text=f"TMO promedio: {tmo:.1f} s", annotation_position="top left")
            fig.add_hline(y=contact_rate, line_dash="dash", line_color="#0056a6", annotation_text=f"Contacto promedio: {contact_rate:.1%}", annotation_position="top right")
            fig.add_vline(x=tmo_q25, line_dash="dot", line_color="#9bb6cc", annotation_text="Q1 TMO", annotation_position="bottom left")
            fig.add_vline(x=tmo_q75, line_dash="dot", line_color="#9bb6cc", annotation_text="Q3 TMO", annotation_position="bottom right")
            fig.add_hline(y=contact_q25, line_dash="dot", line_color="#9bb6cc", annotation_text="Q1 contacto", annotation_position="bottom right")
            fig.add_hline(y=contact_q75, line_dash="dot", line_color="#9bb6cc", annotation_text="Q3 contacto", annotation_position="top left")
            fig.update_layout(
                height=560,
                margin=dict(l=0, r=0, t=90, b=0),
                legend=dict(orientation="h", x=.5, y=-.16, xanchor="center", yanchor="top"),
            )
            fig.update_yaxes(tickformat=".0%", showgrid=True, gridcolor="#edf3f8")
            fig.update_xaxes(showgrid=True, gridcolor="#edf3f8")
            style_chart(fig, 560)
            st.plotly_chart(fig, use_container_width=True, config=PLOT_CONFIG)
            critical = int(campaign_scatter["segment"].eq("Crítica: bajo contacto / alto TMO").sum())
            benchmark = int(campaign_scatter["segment"].eq("Benchmark: alto contacto / bajo TMO").sum())
            potential = int(campaign_scatter["segment"].eq("Potencial: alto contacto / alto TMO").sum())
            st.markdown(
                f'<div class="insight"><strong>Lectura analítica:</strong> {critical} campaña(s) están en zona crítica y deben priorizar reducción de TMO y revisión de horario/script; {potential} tienen buen contacto pero consumen demasiado tiempo; {benchmark} son benchmark para replicar por su combinación de contacto y eficiencia.</div>',
                unsafe_allow_html=True,
            )
        else:
            st.info("No hay campañas con volumen suficiente para comparar TMO y contactabilidad.")
    else:
        st.info("No hay campañas disponibles para comparar TMO y contactabilidad.")

    insight_cards = insight_records(result, total, tmo, contact_rate)
    st.markdown('<div class="section-title">Insights ejecutivos y oportunidades de mejora</div>', unsafe_allow_html=True)
    status_order = {"red": 0, "amber": 1, "green": 2}
    insight_cards.sort(key=lambda card: status_order[card[0]])
    insight_columns = st.columns(3)
    for column, card in zip(insight_columns, insight_cards):
        with column:
            render_insight_card(*card)


if __name__ == "__main__":
    main()